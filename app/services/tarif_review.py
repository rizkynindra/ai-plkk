from __future__ import annotations

import json
import re
import threading
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path

import pymupdf
from openpyxl import Workbook, load_workbook
from rapidfuzz import fuzz, process

from app.config import ROOT_DIR
from app.schemas import TariffReviewDraft
from app.services.ingest.image import ocr_pil_image
from app.services.vision import (
    extract_reference_tariffs_from_image,
    identify_hospital_from_image,
    vision_service_reachable,
)
from app.web import render_app_page

CASE_ID = "emc-alam-sutera"
WORKBOOK_PATH = ROOT_DIR / "master-tarif" / "TARIF_KESEPAKATAN_EMC_ALAM_SUTERA.xlsx"
DATA_DIR = ROOT_DIR / "data" / "tarif-review"
MAX_PDF_BYTES = 40_000_000
RENDER_SCALE = 1.8
IDENTITY_PAGE_LIMIT = 6
HOSPITAL_MATCH_THRESHOLD = 78
PAGE_PATTERN = re.compile(r"hal\.(\d+)", re.IGNORECASE)
REQUIRED_COLUMNS = (
    "kode_master",
    "kategori_bpjs",
    "hospital_name",
    "nama",
    "group_code",
    "price_final",
    "source_page",
    "price_basis",
)

_lock = threading.Lock()
_catalog_cache: dict | None = None
_page_cache: dict[tuple, bytes] = {}
_index_cache: dict[tuple, dict] = {}


def load_catalog() -> dict:
    global _catalog_cache
    if _catalog_cache is None:
        _catalog_cache = _read_catalog()
    return _catalog_cache


def load_draft() -> dict:
    path = _draft_path()
    if not path.exists():
        return {"rows": {}}
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError("Review draft file is unreadable.") from exc
    rows = payload.get("rows") if isinstance(payload, dict) else None
    if not isinstance(rows, dict):
        raise ValueError("Review draft file is unreadable.")
    return {"rows": rows}


def save_draft(draft: TariffReviewDraft) -> dict:
    analysis_rows = pdf_index().get("rows", []) if _pdf_path().exists() else []
    known = {row["kode_master"]: row for row in analysis_rows}
    cleaned: dict[str, dict] = {}
    for kode, item in draft.rows.items():
        source = known.get(kode)
        if source is None:
            raise ValueError(f"Unknown Kode Master: {kode}")
        if item.status == "matches" and item.verified_price != source["price_final"]:
            raise ValueError(f"{kode}: a matching row must keep price_final.")
        if item.status == "corrected" and item.verified_price == source["price_final"]:
            raise ValueError(f"{kode}: a corrected row needs a different price.")
        cleaned[kode] = {
            "verified_price": item.verified_price,
            "status": item.status,
            "note": item.note.strip(),
        }
    payload = {"rows": cleaned}
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    _draft_path().write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def pdf_info() -> dict:
    path = _pdf_path()
    meta = _read_meta()
    if meta is None or not path.exists():
        return {"opened": False}
    return {
        "opened": True,
        "filename": meta["filename"],
        "page_count": meta["page_count"],
        "revision": path.stat().st_mtime_ns,
    }


def store_pdf(payload: bytes, filename: str) -> dict:
    if not payload:
        raise ValueError("Empty file.")
    if len(payload) > MAX_PDF_BYTES:
        raise ValueError("PDF is larger than 40 MB.")
    if not payload.startswith(b"%PDF"):
        raise ValueError("This file is not a PDF.")
    try:
        document = pymupdf.open(stream=payload, filetype="pdf")
    except Exception as exc:
        raise ValueError("This PDF could not be opened.") from exc
    try:
        if document.needs_pass:
            raise ValueError("This PDF is password protected.")
        page_count = document.page_count
    finally:
        document.close()
    if page_count < 1:
        raise ValueError("This PDF has no pages.")
    safe_name = Path(filename).name or "pks.pdf"
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    _pdf_path().write_bytes(payload)
    _meta_path().write_text(
        json.dumps({"filename": safe_name, "page_count": page_count}),
        encoding="utf-8",
    )
    with _lock:
        _page_cache.clear()
        _index_cache.clear()
    return pdf_info()


def render_page_png(page_number: int) -> bytes:
    path = _pdf_path()
    if not path.exists():
        raise FileNotFoundError("Open the PKS PDF first.")
    key = (str(path), path.stat().st_mtime_ns, page_number)
    with _lock:
        cached = _page_cache.get(key)
    if cached is not None:
        return cached
    document = pymupdf.open(path)
    try:
        if page_number < 1 or page_number > document.page_count:
            raise IndexError
        page = document[page_number - 1]
        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(RENDER_SCALE, RENDER_SCALE), alpha=False)
        data = pixmap.tobytes("png")
    finally:
        document.close()
    with _lock:
        _page_cache[key] = data
        if len(_page_cache) > 8:
            _page_cache.pop(next(iter(_page_cache)))
    return data


def pdf_index() -> dict:
    path = _pdf_path()
    if not path.exists():
        raise FileNotFoundError("Open the PKS PDF first.")
    key = (str(path), path.stat().st_mtime_ns)
    with _lock:
        cached = _index_cache.get(key)
    if cached is not None:
        return cached
    catalog = load_catalog()
    document_hospital, identity_method = _identify_document_hospital(path, catalog["hospitals"])
    matched_hospital = _match_hospital(document_hospital, catalog["hospitals"])
    if matched_hospital is None:
        payload = {
            "document_hospital": document_hospital,
            "matched_hospital": None,
            "hospital_status": "unknown" if not document_hospital else "no_reference",
            "identity_method": identity_method,
            "extraction_method": "none",
            "rows": [],
            "hits": {},
            "order": {},
        }
        with _lock:
            _index_cache.clear()
            _index_cache[key] = payload
        return payload

    references = [
        row for row in catalog["rows"] if row["hospital_name"] == matched_hospital
    ]
    rows, hits, extraction_method = _extract_document_tariffs(path, references)
    grouped: dict[int, list[tuple[float, str]]] = {}
    for kode, hit in hits.items():
        grouped.setdefault(hit["pdf_page"], []).append((hit["y"], kode))
    order = {
        str(page_number): [kode for _top, kode in sorted(items)]
        for page_number, items in grouped.items()
    }
    payload = {
        "document_hospital": document_hospital,
        "matched_hospital": matched_hospital,
        "hospital_status": "matched",
        "identity_method": identity_method,
        "extraction_method": extraction_method,
        "rows": rows,
        "hits": hits,
        "order": order,
    }
    with _lock:
        _index_cache.clear()
        _index_cache[key] = payload
    return payload


def marks_for_row(page_number: int, kode: str) -> dict:
    row = _row_by_kode(kode)
    if row is None:
        raise KeyError(kode)
    path = _pdf_path()
    if not path.exists():
        raise FileNotFoundError("Open the PKS PDF first.")
    document = pymupdf.open(path)
    try:
        if page_number < 1 or page_number > document.page_count:
            raise IndexError
        rects, kind = _search_page(document[page_number - 1], row["nama"], row["price_final"])
    finally:
        document.close()
    return {"kind": kind, "rects": rects}


def export_workbook() -> bytes:
    catalog = load_catalog()
    analysis = pdf_index() if _pdf_path().exists() else {"rows": []}
    draft = load_draft()["rows"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reviewed"
    sheet.append(
        [
            "Kode Master",
            "kategori_bpjs",
            "hospital_name",
            "nama",
            "group_code",
            "price_final",
            "reference_price",
            "verified_price",
            "check",
            "note",
        ]
    )
    for row in analysis["rows"]:
        saved = draft.get(row["kode_master"], {})
        sheet.append(
            [
                row["kode_master"],
                row["kategori_bpjs"],
                row["hospital_name"],
                row["nama"],
                row["group_code"] or "",
                row["price_final"],
                row["reference_price"],
                saved.get("verified_price", row["price_final"]),
                saved.get("status", ""),
                saved.get("note", ""),
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _identify_document_hospital(path: Path, candidates: list[str]) -> tuple[str, str]:
    document = pymupdf.open(path)
    page_images: list[bytes] = []
    text_parts: list[str] = []
    try:
        for index in range(min(document.page_count, IDENTITY_PAGE_LIMIT)):
            page = document[index]
            text = page.get_text("text").strip()
            if text:
                text_parts.append(text)
            if index < 3:
                page_images.append(_render_page_bytes(page, 1.0))
    finally:
        document.close()

    text = "\n".join(text_parts)
    hospital = _document_hospital_from_text(text, candidates)
    if hospital:
        return hospital, "pdf_text"

    for image in page_images:
        ocr_text = ocr_pil_image(_image_from_bytes(image))
        hospital = _document_hospital_from_text(ocr_text, candidates)
        if hospital:
            return hospital, "ocr"

    if vision_service_reachable():
        for image in page_images:
            hospital = identify_hospital_from_image(image, candidates)
            if hospital:
                return hospital, "vision"
    return "", "unreadable"


def _document_hospital_from_text(text: str, candidates: list[str]) -> str:
    folded = _normalize_hospital(text)
    if not folded:
        return ""
    for candidate in candidates:
        normalized = _normalize_hospital(candidate)
        if normalized and normalized in folded:
            return candidate
    for line in text.splitlines():
        clean = " ".join(line.split()).strip(" .,:;-")
        match = re.search(r"\b(?:rumah\s+sakit|rs)\s+([a-z0-9][a-z0-9 .&'-]{2,60})", clean, re.I)
        if match:
            return f"RS {match.group(1).strip()}"
    return ""


def _match_hospital(document_hospital: str, candidates: list[str]) -> str | None:
    if not document_hospital:
        return None
    normalized = _normalize_hospital(document_hospital)
    choices = {_normalize_hospital(candidate): candidate for candidate in candidates}
    result = process.extractOne(normalized, choices.keys(), scorer=fuzz.WRatio)
    if result is None or result[1] < HOSPITAL_MATCH_THRESHOLD:
        return None
    return choices[result[0]]


def _normalize_hospital(value: str) -> str:
    text = value.lower()
    text = re.sub(r"\brumah\s+sakit\b|\brs\b", " ", text)
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _extract_document_tariffs(
    path: Path,
    references: list[dict],
) -> tuple[list[dict], dict[str, dict], str]:
    text_rows, text_hits = _extract_from_text_pdf(path, references)
    if text_rows:
        return text_rows, text_hits, "pdf_text"
    if not vision_service_reachable():
        return [], {}, "vision_unavailable"
    return (*_extract_from_scanned_pdf(path, references), "vision")


def _extract_from_text_pdf(
    path: Path,
    references: list[dict],
) -> tuple[list[dict], dict[str, dict]]:
    prepared = [_prepared_row(row) for row in references]
    by_code = {row["kode_master"]: row for row in references}
    extracted: list[dict] = []
    hits: dict[str, dict] = {}
    document = pymupdf.open(path)
    try:
        for page_number, page in enumerate(document, start=1):
            folded = " ".join(page.get_text("text").split()).lower()
            if not folded:
                continue
            for item in prepared:
                if item["kode"] in hits:
                    continue
                term = _matching_term(folded, item)
                if term is None:
                    continue
                rects = _rects(page, term)
                if not rects:
                    continue
                price = _price_near_rect(page, rects[0])
                if price is None:
                    continue
                reference = by_code[item["kode"]]
                extracted.append(_extracted_row(reference, price, page_number, term))
                hits[item["kode"]] = {
                    "pdf_page": page_number,
                    "y": rects[0]["y"],
                }
    finally:
        document.close()
    return extracted, hits


def _extract_from_scanned_pdf(
    path: Path,
    references: list[dict],
) -> tuple[list[dict], dict[str, dict]]:
    by_page: dict[int, list[dict]] = {}
    for row in references:
        if row["page"]:
            by_page.setdefault(row["page"], []).append(row)

    def extract_page(page_number: int, page_rows: list[dict]) -> tuple[int, list[dict]]:
        document = pymupdf.open(path)
        try:
            if page_number < 1 or page_number > document.page_count:
                return page_number, []
            image = _render_page_bytes(document[page_number - 1], 1.5)
        finally:
            document.close()
        return page_number, extract_reference_tariffs_from_image(image, page_rows)

    results: list[tuple[int, list[dict]]] = []
    with ThreadPoolExecutor(max_workers=min(4, len(by_page) or 1)) as executor:
        futures = [
            executor.submit(extract_page, page_number, rows)
            for page_number, rows in by_page.items()
        ]
        results = [future.result() for future in futures]

    references_by_code = {row["kode_master"]: row for row in references}
    extracted: list[dict] = []
    hits: dict[str, dict] = {}
    for page_number, items in results:
        for position, item in enumerate(items):
            if not isinstance(item, dict):
                continue
            kode = str(item.get("kode_master") or "").strip()
            reference = references_by_code.get(kode)
            price = item.get("price")
            if reference is None or not isinstance(price, (int, float)) or price < 0:
                continue
            if kode in hits:
                continue
            printed_name = str(item.get("printed_name") or reference["nama"]).strip()
            extracted.append(_extracted_row(reference, int(price), page_number, printed_name))
            hits[kode] = {"pdf_page": page_number, "y": position / max(len(items), 1)}
    return extracted, hits


def _extracted_row(reference: dict, price: int, page: int, printed_name: str) -> dict:
    row = dict(reference)
    row["reference_price"] = reference["price_final"]
    row["price_final"] = price
    row["printed_name"] = printed_name
    row["pdf_page"] = page
    row["price_matches_reference"] = price == reference["price_final"]
    return row


def _price_near_rect(page: pymupdf.Page, normalized_rect: dict) -> int | None:
    page_width = float(page.rect.width) or 1
    page_height = float(page.rect.height) or 1
    y0 = normalized_rect["y"] * page_height
    y1 = (normalized_rect["y"] + normalized_rect["h"]) * page_height
    amounts: list[tuple[float, int]] = []
    for word in page.get_text("words"):
        x0, word_y0, _x1, word_y1, text = word[:5]
        if word_y1 < y0 - 4 or word_y0 > y1 + 4:
            continue
        parsed = _parse_price(text)
        if parsed is not None:
            amounts.append((x0 / page_width, parsed))
    if not amounts:
        return None
    return max(amounts, key=lambda item: item[0])[1]


def _parse_price(value: str) -> int | None:
    compact = re.sub(r"[^\d]", "", value)
    if not compact:
        return None
    amount = int(compact)
    return amount if amount >= 1000 else None


def _render_page_bytes(page: pymupdf.Page, scale: float) -> bytes:
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
    return pixmap.tobytes("png")


def _image_from_bytes(data: bytes):
    from PIL import Image

    return Image.open(BytesIO(data))


def _read_catalog() -> dict:
    if not WORKBOOK_PATH.exists():
        raise FileNotFoundError("Tariff workbook was not found.")
    workbook = load_workbook(WORKBOOK_PATH, read_only=True, data_only=True)
    try:
        period = _period(workbook)
        sheet = workbook["Tariffs"]
        rows_iter = sheet.iter_rows(values_only=True)
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows_iter)]
        index = {name: header.index(name) for name in REQUIRED_COLUMNS}
        rows = []
        seen: set[str] = set()
        for raw in rows_iter:
            if not any(value not in (None, "") for value in raw):
                continue
            kode = str(raw[index["kode_master"]] or "").strip()
            if not kode or kode in seen:
                raise ValueError(f"Tariff workbook has a missing or duplicate Kode Master: {kode}")
            seen.add(kode)
            source_page = str(raw[index["source_page"]] or "").strip()
            group = raw[index["group_code"]]
            rows.append(
                {
                    "kode_master": kode,
                    "kategori_bpjs": str(raw[index["kategori_bpjs"]] or "").strip(),
                    "hospital_name": str(raw[index["hospital_name"]] or "").strip(),
                    "nama": str(raw[index["nama"]] or "").strip(),
                    "group_code": str(group).strip() if group not in (None, "") else None,
                    "price_final": int(raw[index["price_final"]]),
                    "source_page": source_page,
                    "page": _printed_page(source_page),
                    "price_basis": str(raw[index["price_basis"]] or "").strip(),
                }
            )
    finally:
        workbook.close()
    if not rows:
        raise ValueError("Tariff workbook has no rows.")
    pages = sorted({row["page"] for row in rows if row["page"]})
    hospitals = sorted({row["hospital_name"] for row in rows if row["hospital_name"]})
    return {
        "case_id": CASE_ID,
        "workbook": WORKBOOK_PATH.name,
        "hospital_name": rows[0]["hospital_name"],
        "hospitals": hospitals,
        "period": period,
        "row_count": len(rows),
        "printed_pages": pages,
        "rows": rows,
    }


def _period(workbook) -> str | None:
    if "Aturan Tarif" not in workbook.sheetnames:
        return None
    sheet = workbook["Aturan Tarif"]
    rows = sheet.iter_rows(values_only=True)
    header = [str(cell).strip() if cell is not None else "" for cell in next(rows)]
    if "topik" not in header or "angka" not in header:
        return None
    topic_index = header.index("topik")
    number_index = header.index("angka")
    for row in rows:
        topic = str(row[topic_index] or "")
        if "masa berlaku" not in topic.lower():
            continue
        angka = str(row[number_index] or "").strip()
        return angka or None
    return None


def _printed_page(source_page: str) -> int | None:
    match = PAGE_PATTERN.search(source_page)
    if match is None:
        return None
    return int(match.group(1))


def _row_by_kode(kode: str) -> dict | None:
    for row in load_catalog()["rows"]:
        if row["kode_master"] == kode:
            return row
    return None


def _search_page(page: pymupdf.Page, nama: str, price: int) -> tuple[list[dict], str]:
    for term in _name_terms(nama):
        rects = _rects(page, term)
        if rects:
            return rects[:6], "name"
    dotted = f"{int(price):,}".replace(",", ".")
    price_rects = _rects(page, dotted)[:4]
    if price_rects:
        return price_rects, "price"
    return [], "none"


def _prepared_row(row: dict) -> dict:
    terms = []
    if row.get("group_code"):
        terms.append(str(row["group_code"]))
    terms.extend(_name_terms(row["nama"]))
    return {
        "kode": row["kode_master"],
        "terms": terms,
        "price": f"{int(row['price_final']):,}".replace(",", "."),
        "grouped": bool(row.get("group_code")),
    }


def _matching_term(folded_text: str, item: dict) -> str | None:
    if item["grouped"]:
        code = item["terms"][0]
        if code.lower() in folded_text and item["price"] in folded_text:
            return code
        return None
    for term in item["terms"]:
        if " ".join(term.split()).lower() in folded_text:
            return term
    return None


def _name_terms(nama: str) -> list[str]:
    cleaned = " ".join(nama.split())
    head = cleaned.split(" (", 1)[0].strip()
    terms = []
    if len(head) >= 4:
        terms.append(head[:80])
    if cleaned != head and 4 <= len(cleaned) <= 80:
        terms.append(cleaned)
    return terms


def _rects(page: pymupdf.Page, term: str) -> list[dict]:
    width = float(page.rect.width) or 1
    height = float(page.rect.height) or 1
    found = []
    for rect in page.search_for(term):
        found.append(
            {
                "x": rect.x0 / width,
                "y": rect.y0 / height,
                "w": rect.width / width,
                "h": rect.height / height,
            }
        )
    return found


def _draft_path() -> Path:
    return DATA_DIR / f"{CASE_ID}-draft.json"


def _pdf_path() -> Path:
    return DATA_DIR / f"{CASE_ID}.pdf"


def _meta_path() -> Path:
    return DATA_DIR / f"{CASE_ID}-pdf.json"


def _read_meta() -> dict | None:
    path = _meta_path()
    if not path.exists():
        return None
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        return None
    if "filename" not in payload or "page_count" not in payload:
        return None
    return payload


REVIEW_BODY = """
      <section class="review-desk" id="review-desk">
        <p id="review-status" class="status" role="status">Reading TARIF_KESEPAKATAN_EMC_ALAM_SUTERA.xlsx…</p>
        <div id="review-error" class="status is-error" hidden>
          <p id="review-error-text"></p>
          <button id="retry-load" class="button secondary" type="button">Read the workbook again</button>
        </div>
        <dl id="doc-facts" class="doc-facts" hidden>
          <div><dt>Workbook</dt><dd id="fact-workbook"></dd></div>
          <div><dt>Document hospital</dt><dd id="fact-hospital"></dd></div>
          <div><dt>Extracted rows</dt><dd id="fact-rows"></dd></div>
          <div><dt>Reference match</dt><dd id="fact-pages"></dd></div>
          <div id="fact-period" hidden><dt>Period</dt><dd></dd></div>
        </dl>
        <div class="review-split">
          <section class="pdf-pane" aria-labelledby="pdf-title">
            <div class="pane-bar">
              <h2 id="pdf-title">Printed page</h2>
            </div>
            <div class="pdf-tools">
              <label class="pdf-open">Open PDF
                <input id="pdf-input" type="file" accept="application/pdf,.pdf">
              </label>
              <button id="page-prev" class="button secondary" type="button">Previous</button>
              <label class="page-field">Page
                <input id="page-input" inputmode="numeric" value="1">
              </label>
              <span id="page-total"></span>
              <button id="page-next" class="button secondary" type="button">Next</button>
              <button id="zoom-out" class="button secondary" type="button">Smaller</button>
              <button id="zoom-in" class="button secondary" type="button">Larger</button>
            </div>
            <p id="pdf-status" class="status" role="status">Open the PKS PDF. Rows on the right follow the page in view.</p>
            <div id="pdf-stage" class="pdf-stage">
              <p id="pdf-empty" class="pdf-empty">Open the PKS PDF. The rows beside it list only the lines found on the page you are viewing.</p>
              <div id="pdf-frame" class="pdf-frame" hidden>
                <img id="pdf-image" alt="">
                <div id="pdf-marks" aria-hidden="true"></div>
              </div>
            </div>
            <section id="row-detail" class="row-detail" hidden>
              <h3 id="detail-nama" class="detail-title"></h3>
              <p id="detail-meta"></p>
              <p id="detail-check"></p>
              <dl class="detail-list">
                <div><dt>Extracted price</dt><dd id="detail-price"></dd></div>
                <div><dt>Reference price</dt><dd id="detail-reference-price"></dd></div>
                <div><dt>Cited page</dt><dd id="detail-page"></dd></div>
              </dl>
              <p id="detail-basis" class="basis"></p>
              <div class="form-grid">
                <label>Printed amount
                  <input id="verified-price" inputmode="numeric">
                </label>
                <label class="wide">Note
                  <textarea id="review-note" maxlength="500"></textarea>
                </label>
              </div>
              <div class="actions">
                <button id="show-cited" class="button secondary" type="button">Show on the PDF</button>
                <button id="mark-match" class="button secondary" type="button">Printed price matches</button>
                <button id="mark-corrected" class="button" type="button">Use corrected price</button>
                <button id="clear-check" class="button secondary" type="button">Clear this check</button>
              </div>
            </section>
          </section>
          <section class="extract-pane" aria-labelledby="extract-title">
            <div class="pane-bar">
              <h2 id="extract-title">Extracted rows</h2>
              <p id="row-count"></p>
            </div>
            <div class="review-filters">
              <label>Search
                <input id="row-search" type="search" placeholder="Kode or nama">
              </label>
              <label>kategori_bpjs
                <select id="kategori-filter">
                  <option value="">All categories</option>
                </select>
              </label>
              <label>Check
                <select id="status-filter">
                  <option value="all">All rows</option>
                  <option value="unchecked">Still unchecked</option>
                  <option value="matches">Matches printed page</option>
                  <option value="corrected">Corrected</option>
                </select>
              </label>
              <label>Scope
                <select id="page-scope">
                  <option value="page">On this PDF page</option>
                  <option value="all">All extracted rows</option>
                </select>
              </label>
            </div>
            <div class="table-wrap review-table">
              <table aria-labelledby="extract-title">
                <thead>
                  <tr>
                    <th>Kode Master</th>
                    <th>nama</th>
                    <th>price_final</th>
                    <th>Reference price</th>
                    <th>Check</th>
                    <th>kategori_bpjs</th>
                    <th>group_code</th>
                    <th>hospital_name</th>
                  </tr>
                </thead>
                <tbody id="extract-body"></tbody>
              </table>
            </div>
            <p id="filter-empty" class="status is-empty" hidden>No rows match this filter. Clear the search or choose All categories.</p>
            <p class="key-hint">Arrow keys move through the visible rows.</p>
          </section>
        </div>
        <section class="panel decision">
          <p id="decision-summary" class="status" role="status"></p>
          <div class="actions">
            <button id="save-draft" class="button" type="button">Save review draft</button>
            <a id="download-review" class="button secondary" href="/api/tarif-review/emc-alam-sutera/export">Download reviewed rows</a>
          </div>
          <p id="save-status" class="status" role="status"></p>
        </section>
      </section>
"""

REVIEW_PAGE = render_app_page(
    title="Review extracted tariff, EPLKK",
    heading="Review extracted tariff",
    lede="Open a hospital tariff PDF. The hospital is read from the document, then its printed rows are extracted using the matching workbook references.",
    active="review",
    body=REVIEW_BODY,
    scripts=["/static/nav.js", "/static/tarif-review.js"],
    shell_class="shell shell-wide",
    body_class="review-body",
)
