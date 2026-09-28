from __future__ import annotations

from io import BytesIO

from docx import Document

from app.schemas import DocumentContent, ExtractedTable
from app.services.ingest import cell_str


def ingest_docx(data: bytes, filename: str) -> DocumentContent:
    document = Document(BytesIO(data))
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tables: list[ExtractedTable] = []
    for table in document.tables:
        rows = [[cell_str(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        tables.append(ExtractedTable(headers=rows[0], rows=rows[1:]))
    return DocumentContent(
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        source="word",
        pages_text=paragraphs,
        tables=tables,
    )
