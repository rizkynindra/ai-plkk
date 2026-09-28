from __future__ import annotations

from copy import copy
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from app.config import settings
from app.schemas import CostLine

HEADER_ROW = 1
FIRST_DATA_ROW = 2
TEMPLATE_SHEET = "Template"

COL_NO = 1
COL_ITEM = 2
COL_QTY = 3
COL_TARIFF = 4
COL_COST = 5


def build_template_workbook(lines: list[CostLine], template_path: Path | None = None) -> bytes:
    path = template_path or settings.template_path
    workbook = load_workbook(path)
    sheet = workbook[TEMPLATE_SHEET] if TEMPLATE_SHEET in workbook.sheetnames else workbook.active
    styles = [_cell_style(sheet.cell(FIRST_DATA_ROW, col)) for col in range(COL_NO, COL_COST + 1)]
    _clear_data_rows(sheet)
    for offset, line in enumerate(lines):
        row_idx = FIRST_DATA_ROW + offset
        values = (line.sequence_no, line.item_name, line.quantity, line.tariff, line.discounted_cost)
        for col, (value, style) in enumerate(zip(values, styles), start=COL_NO):
            cell = sheet.cell(row_idx, col, value)
            _apply_style(cell, style)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _clear_data_rows(sheet: Worksheet) -> None:
    if sheet.max_row and sheet.max_row >= FIRST_DATA_ROW:
        sheet.delete_rows(FIRST_DATA_ROW, sheet.max_row - HEADER_ROW)


def _cell_style(cell: Cell) -> dict:
    return {
        "font": copy(cell.font),
        "border": copy(cell.border),
        "fill": copy(cell.fill),
        "number_format": cell.number_format,
        "alignment": copy(cell.alignment),
    }


def _apply_style(cell: Cell, style: dict) -> None:
    cell.font = copy(style["font"])
    cell.border = copy(style["border"])
    cell.fill = copy(style["fill"])
    cell.number_format = style["number_format"]
    cell.alignment = copy(style["alignment"])
