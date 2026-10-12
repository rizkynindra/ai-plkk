from io import BytesIO

import pymupdf
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from app.services.tarif_review import load_catalog

client = TestClient(app)
CASE = "/api/tarif-review/emc-alam-sutera"


def test_review_page_lists_extracted_columns() -> None:
    response = client.get("/review")
    assert response.status_code == 200
    page = response.text
    assert "Review extracted tariff" in page
    assert "Kode Master" in page
    assert "kategori_bpjs" in page
    assert "hospital_name" in page
    assert "price_final" in page
    assert "Open PDF" in page
    assert "Printed price matches" in page
    assert "On this PDF page" in page
    assert page.find('class="pdf-pane"') < page.find('class="extract-pane"')
    assert "Use corrected price" in page
    assert 'href="/api/tarif-review/emc-alam-sutera/export"' in page
    assert "\u2014" not in page
    assert "Siti Nurhaliza" not in page
    home = client.get("/")
    assert 'href="/review"' in home.text


def test_catalog_reads_emc_workbook() -> None:
    response = client.get(CASE)
    assert response.status_code == 200
    body = response.json()
    assert body["workbook"] == "TARIF_KESEPAKATAN_EMC_ALAM_SUTERA.xlsx"
    assert body["hospital_name"] == "RS EMC Alam Sutera"
    assert body["row_count"] == 165
    assert body["printed_pages"] == [25, 26, 27, 28]
    assert "2026" in body["period"]
    assert len({row["kode_master"] for row in body["rows"]}) == 165
    sample = next(row for row in body["rows"] if row["kode_master"] == "EMCRAD0019")
    assert sample["kategori_bpjs"] == "RONTGEN"
    assert sample["nama"].startswith("abdomen atas")
    assert sample["group_code"] is None
    assert sample["price_final"] == 750000
    assert sample["page"] == 26
    grouped = next(row for row in body["rows"] if row["group_code"] == "KECIL")
    assert grouped["price_final"] == 10000000


def test_draft_roundtrip_and_export(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.services.tarif_review.DATA_DIR", tmp_path)
    rows = []
    for code in ("EMCRAD0019", "EMCTOP0006"):
        source = next(row for row in load_catalog()["rows"] if row["kode_master"] == code)
        extracted = dict(source)
        extracted["reference_price"] = source["price_final"]
        rows.append(extracted)
    (tmp_path / "emc-alam-sutera.pdf").write_bytes(b"%PDF-test")
    monkeypatch.setattr("app.services.tarif_review.pdf_index", lambda: {"rows": rows})
    empty = client.get(CASE + "/draft")
    assert empty.status_code == 200
    assert empty.json() == {"rows": {}}

    saved = client.put(
        CASE + "/draft",
        json={
            "rows": {
                "EMCRAD0019": {"verified_price": 700000, "status": "corrected", "note": "Printed lower"},
                "EMCTOP0006": {"verified_price": 10000000, "status": "matches", "note": ""},
            }
        },
    )
    assert saved.status_code == 200
    assert saved.json()["rows"]["EMCRAD0019"]["verified_price"] == 700000

    rejected = client.put(
        CASE + "/draft",
        json={"rows": {"EMCRAD0019": {"verified_price": 750000, "status": "corrected", "note": ""}}},
    )
    assert rejected.status_code == 422

    unknown = client.put(
        CASE + "/draft",
        json={"rows": {"NOPE": {"verified_price": 1, "status": "corrected", "note": ""}}},
    )
    assert unknown.status_code == 422

    export = client.get(CASE + "/export")
    assert export.status_code == 200
    workbook = load_workbook(BytesIO(export.content), read_only=True, data_only=True)
    sheet = workbook["Reviewed"]
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    header = rows[0]
    assert header[:6] == ("Kode Master", "kategori_bpjs", "hospital_name", "nama", "group_code", "price_final")
    corrected = next(row for row in rows[1:] if row[0] == "EMCRAD0019")
    assert corrected[6] == 750000
    assert corrected[7] == 700000
    assert corrected[8] == "corrected"


def test_pdf_page_and_marks(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.services.tarif_review.DATA_DIR", tmp_path)
    info = client.get(CASE + "/pdf")
    assert info.json()["opened"] is False

    rejected = client.post(CASE + "/pdf", files={"file": ("notes.txt", b"hello", "text/plain")})
    assert rejected.status_code == 422

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "abdomen atas/bawah/lengkap")
    page.insert_text((72, 110), "750.000")
    payload = document.tobytes()
    document.close()

    uploaded = client.post(
        CASE + "/pdf",
        files={"file": ("PKS EMC.pdf", payload, "application/pdf")},
    )
    assert uploaded.status_code == 200
    assert uploaded.json()["opened"] is True
    assert uploaded.json()["page_count"] == 1
    assert uploaded.json()["filename"] == "PKS EMC.pdf"

    image = client.get(CASE + "/pdf/pages/1")
    assert image.status_code == 200
    assert image.content.startswith(b"\x89PNG")

    missing = client.get(CASE + "/pdf/pages/9")
    assert missing.status_code == 404

    marks = client.get(CASE + "/pdf/pages/1/marks", params={"kode": "EMCRAD0019"})
    assert marks.status_code == 200
    body = marks.json()
    assert body["kind"] == "name"
    assert body["rects"]
    assert 0 <= body["rects"][0]["x"] <= 1
    assert body["rects"][0]["w"] > 0


def test_pdf_index_uses_the_page_that_contains_the_name(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.services.tarif_review.DATA_DIR", tmp_path)
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "RS EMC Alam Sutera")
    tariff = document.new_page()
    tariff.insert_text((72, 720), "abdomen atas/bawah/lengkap")
    tariff.insert_text((400, 720), "750.000")
    payload = document.tobytes()
    document.close()
    uploaded = client.post(CASE + "/pdf", files={"file": ("pks.pdf", payload, "application/pdf")})
    assert uploaded.status_code == 200
    index = client.get(CASE + "/pdf/index")
    assert index.status_code == 200
    hit = index.json()["hits"]["EMCRAD0019"]
    assert hit["pdf_page"] == 2
    assert index.json()["order"]["2"][0] == "EMCRAD0019"
    assert index.json()["document_hospital"] == "RS EMC Alam Sutera"
    assert index.json()["rows"][0]["price_final"] == 750000


def test_pdf_hospital_without_reference_returns_no_emc_rows(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.services.tarif_review.DATA_DIR", tmp_path)
    document = pymupdf.open()
    document.new_page().insert_text((72, 72), "RS Columbia Asia BSD")
    payload = document.tobytes()
    document.close()
    uploaded = client.post(CASE + "/pdf", files={"file": ("document.pdf", payload, "application/pdf")})
    assert uploaded.status_code == 200
    result = client.get(CASE + "/pdf/index")
    assert result.status_code == 200
    assert result.json()["document_hospital"] == "RS Columbia Asia BSD"
    assert result.json()["hospital_status"] == "no_reference"
    assert result.json()["matched_hospital"] is None
    assert result.json()["rows"] == []


def test_scanned_pdf_uses_vision_after_hospital_match(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr("app.services.tarif_review.DATA_DIR", tmp_path)
    monkeypatch.setattr("app.services.tarif_review.vision_service_reachable", lambda: True)
    monkeypatch.setattr(
        "app.services.tarif_review.identify_hospital_from_image",
        lambda _image, _candidates: "RS EMC Alam Sutera",
    )

    def fake_extract(_image, references):
        row = references[0]
        return [
            {
                "kode_master": row["kode_master"],
                "printed_name": row["nama"],
                "price": row["price_final"] + 1000,
            }
        ]

    monkeypatch.setattr(
        "app.services.tarif_review.extract_reference_tariffs_from_image",
        fake_extract,
    )
    document = pymupdf.open()
    for _ in range(28):
        document.new_page()
    payload = document.tobytes()
    document.close()
    uploaded = client.post(CASE + "/pdf", files={"file": ("scan.pdf", payload, "application/pdf")})
    assert uploaded.status_code == 200
    result = client.get(CASE + "/pdf/index")
    assert result.status_code == 200
    body = result.json()
    assert body["matched_hospital"] == "RS EMC Alam Sutera"
    assert body["identity_method"] == "vision"
    assert body["extraction_method"] == "vision"
    assert len(body["rows"]) == 4
    assert body["rows"][0]["price_final"] == body["rows"][0]["reference_price"] + 1000
