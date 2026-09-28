from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.convert import NotClaimInvoiceError, convert_document
from app.services.ingest.router import guess_filename
from app.services.money import parse_idr_amount
from app.services.vision import (
    _client_chat_endpoint,
    _is_summary_name,
    _next_chat_endpoint,
    _ollama_opener,
    _ollama_vision_json,
    _prefer_detail_lines,
    _reconcile_unit_and_total,
    _validated_lines,
    looks_like_non_billing_header,
)

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


def test_guess_filename_from_magic_bytes() -> None:
    assert guess_filename((SAMPLES / "Template_Biaya_sample.xlsx").read_bytes()) == "document.xlsx"
    assert guess_filename((SAMPLES / "Template_Biaya_sample.pdf").read_bytes()) == "document.pdf"


def test_parse_idr_grouped_and_plain() -> None:
    assert parse_idr_amount("Rp 1.000.000") == 1_000_000
    assert parse_idr_amount("100000") == 100_000
    assert parse_idr_amount("50.000,50") == 50000.5
    assert parse_idr_amount("HOMECARE") is None


def test_convert_sample_xlsx_as_is() -> None:
    data = (SAMPLES / "Template_Biaya_sample.xlsx").read_bytes()
    preview, workbook = convert_document(data, "Template_Biaya_sample.xlsx")
    assert preview.classification.is_claim_invoice
    assert [line.item_name for line in preview.lines] == ["BIAYA ADMINISTRASI", "HOMECARE"]
    assert preview.lines[0].quantity == 1
    assert preview.lines[0].tariff == 100_000
    assert preview.lines[1].discounted_cost == 100_000

    sheet = load_workbook(BytesIO(workbook)).active
    assert sheet["B2"].value == "BIAYA ADMINISTRASI"
    assert sheet["C3"].value == 2
    assert sheet["E3"].value == 100_000


def test_convert_sample_pdf() -> None:
    data = (SAMPLES / "Template_Biaya_sample.pdf").read_bytes()
    preview, _workbook = convert_document(data, "Template_Biaya_sample.pdf")
    names = [line.item_name for line in preview.lines]
    assert "BIAYA ADMINISTRASI" in names
    assert "HOMECARE" in names


def test_convert_docx_table() -> None:
    from io import BytesIO as Buffer

    from docx import Document

    document = Document()
    table = document.add_table(rows=3, cols=5)
    table.rows[0].cells[0].text = "No"
    table.rows[0].cells[1].text = "Item Biaya"
    table.rows[0].cells[2].text = "Jumlah"
    table.rows[0].cells[3].text = "Tarif"
    table.rows[0].cells[4].text = "Biaya (setelah dikurangi diskon)"
    table.rows[1].cells[0].text = "1"
    table.rows[1].cells[1].text = "BIAYA ADMINISTRASI"
    table.rows[1].cells[2].text = "1"
    table.rows[1].cells[3].text = "100000"
    table.rows[1].cells[4].text = "100000"
    table.rows[2].cells[0].text = "2"
    table.rows[2].cells[1].text = "HOMECARE"
    table.rows[2].cells[2].text = "2"
    table.rows[2].cells[3].text = "50000"
    table.rows[2].cells[4].text = "100000"
    buffer = Buffer()
    document.save(buffer)
    preview, _workbook = convert_document(buffer.getvalue(), "claim.docx")
    assert [line.item_name for line in preview.lines] == ["BIAYA ADMINISTRASI", "HOMECARE"]


def test_convert_csv_table() -> None:
    csv = (
        "No,Item Biaya,Jumlah,Tarif,Biaya (setelah dikurangi diskon)\n"
        "1,BIAYA ADMINISTRASI,1,100000,100000\n"
        "2,HOMECARE,2,50000,100000\n"
    ).encode()
    preview, _workbook = convert_document(csv, "claim.csv")
    assert [line.item_name for line in preview.lines] == ["BIAYA ADMINISTRASI", "HOMECARE"]


def test_reject_iks_catalog() -> None:
    data = (SAMPLES / "Daftar_Biaya_IKS_20260917_022547.xlsx").read_bytes()
    with pytest.raises(NotClaimInvoiceError):
        convert_document(data, "Daftar_Biaya_IKS.xlsx")


def test_reject_unknown_extension() -> None:
    from app.services.ingest.router import UnsupportedDocumentError

    with pytest.raises(UnsupportedDocumentError):
        convert_document(b"hello", "notes.txt")


def test_vision_lines_require_exact_amount_evidence() -> None:
    items = [
        {
            "item_name": "Kamar Perawatan",
            "quantity_text": "2",
            "tariff_text": "Rp 500.000",
            "discounted_cost_text": "Rp 1.000.000",
            "evidence": "Kamar Perawatan 2 Rp 500.000 Rp 1.000.000",
        },
        {
            "item_name": "Invented Cost",
            "quantity_text": "1",
            "tariff_text": "250000",
            "discounted_cost_text": "250000",
            "evidence": "Invented Cost",
        },
    ]

    lines = _validated_lines(items, page_number=12)

    assert len(lines) == 1
    assert lines[0].item_name == "Kamar Perawatan"
    assert lines[0].source_page == 12
    assert lines[0].tariff == 500_000
    assert lines[0].discounted_cost == 1_000_000


def test_reconcile_when_tariff_copied_from_line_total() -> None:
    tariff, cost = _reconcile_unit_and_total(3, 82479, 82479)
    assert cost == 82479
    assert abs(tariff - 27493) < 1


def test_skip_kuitansi_grand_total_row() -> None:
    items = [
        {
            "item_name": "Tagihan Rawat Inap a.n ASEP MULJANAK",
            "quantity_text": "1",
            "tariff_text": "Rp 23.559.342",
            "discounted_cost_text": "Rp 23.559.342",
            "evidence": "Tagihan Rawat Inap a.n ASEP MULJANAK Rp 23.559.342",
        }
    ]
    assert _validated_lines(items, page_number=3) == []
    assert _is_summary_name("Tagihan Rawat Inap a.n ASEP MULJANAK")
    assert not _is_summary_name("Prosogan 30 mg SOL")


def test_prefer_detail_lines_drops_instalasi_summary() -> None:
    from app.schemas import CostLine

    summary = [
        CostLine(sequence_no=1, item_name="INSTALASI RAWAT INAP", quantity=1, tariff=9_829_965, discounted_cost=9_829_965, source_page=4),
        CostLine(sequence_no=2, item_name="Biaya Administrasi", quantity=1, tariff=686_194, discounted_cost=686_194, source_page=4),
    ]
    details = [
        CostLine(sequence_no=i, item_name=f"Obat {i}", quantity=1, tariff=1000, discounted_cost=1000, source_page=5)
        for i in range(1, 10)
    ]
    kept = _prefer_detail_lines(summary + details)
    names = [line.item_name for line in kept]
    assert "INSTALASI RAWAT INAP" not in names
    assert "Biaya Administrasi" in names
    assert len(kept) == 10


def test_skip_non_billing_bundle_headers() -> None:
    assert looks_like_non_billing_header("Dokumen: KARTU PESERTA")
    assert looks_like_non_billing_header("Dokumen: RINGKASAN PULANG")
    assert not looks_like_non_billing_header("Dokumen: DOKUMEN KELENGKAPAN PERSYARATAN")
    assert not looks_like_non_billing_header("")


def test_ollama_chat_endpoint_strips_openai_v1_suffix() -> None:
    assert _next_chat_endpoint().endswith("/api/chat")
    assert "/v1/api/chat" not in _next_chat_endpoint()


def test_ollama_chat_endpoint_rewrites_unspecified_bind_address() -> None:
    assert _client_chat_endpoint("http://0.0.0.0:11434/v1") == "http://127.0.0.1:11434/api/chat"
    assert _client_chat_endpoint("http://[::]:11434/v1") == "http://127.0.0.1:11434/api/chat"
    assert _client_chat_endpoint("http://127.0.0.1:11434/v1") == "http://127.0.0.1:11434/api/chat"


def test_ollama_opener_ignores_env_proxies() -> None:
    handlers = [type(handler).__name__ for handler in _ollama_opener.handlers]
    assert "ProxyHandler" in handlers
    proxy_handler = next(
        handler for handler in _ollama_opener.handlers if type(handler).__name__ == "ProxyHandler"
    )
    assert proxy_handler.proxies == {}


def test_ollama_vision_json_returns_none_on_http_403(monkeypatch) -> None:
    from io import BytesIO
    from urllib.error import HTTPError

    def boom(_request, timeout=0):
        raise HTTPError(
            "http://127.0.0.1:11434/api/chat",
            403,
            "Forbidden",
            hdrs=None,
            fp=BytesIO(b"Forbidden"),
        )

    monkeypatch.setattr("app.services.vision._ollama_opener.open", boom)
    assert _ollama_vision_json(b"jpeg", "prompt") is None
